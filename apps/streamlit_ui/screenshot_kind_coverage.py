"""One-off Playwright verification that the Secondary/CCA example rows
in the blank template AND the demo dataset actually work end-to-end.
Run against a live `streamlit run app.py` server on :8765 with an EMPTY
database, same convention as screenshot_phase5.py.
"""
import os

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

OUT = "/tmp/itap_kind_coverage_screens"
os.makedirs(OUT, exist_ok=True)
BASE = "http://localhost:8765"

with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page(viewport={"width": 1400, "height": 1000})
    page.goto(BASE)
    page.wait_for_timeout(2500)

    # Demo dataset download button lives on the empty-deployment welcome
    # screen, before any Admin exists.
    with page.expect_download() as dl_info2:
        page.get_by_text("Download demo dataset (.xlsx)").click()
    demo_path = f"{OUT}/demo.xlsx"
    dl_info2.value.save_as(demo_path)

    demo_wb = load_workbook(demo_path)
    demo_ws = demo_wb["Assignments"]
    demo_kinds = [r[2] for r in demo_ws.iter_rows(min_row=2, values_only=True)]
    print("Demo dataset Assignments kinds:", {k: demo_kinds.count(k) for k in set(demo_kinds)})
    assert "secondary" in demo_kinds
    assert "cca" in demo_kinds

    # First Admin, so we can reach Bulk Setup.
    page.get_by_text("Add a new person (ITAP Admin setup)").click()
    page.wait_for_timeout(300)
    page.get_by_label("Admin name").fill("Priya")
    page.get_by_text("Create ITAP Admin").click()
    page.wait_for_timeout(1500)

    page.get_by_role("button", name="Priya", exact=True).click()
    page.wait_for_timeout(1200)

    page.get_by_role("tab", name="Bulk Setup").click()
    page.wait_for_timeout(800)

    # -- Blank template download: confirm Secondary/CCA rows are in it. --
    with page.expect_download() as dl_info:
        page.get_by_text("Download template (.xlsx)").click()
    template_path = f"{OUT}/template.xlsx"
    dl_info.value.save_as(template_path)

    wb = load_workbook(template_path)
    ws = wb["Assignments"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    kinds = [r[2] for r in rows]
    print("Template Assignments kinds:", kinds)
    assert "secondary" in kinds, "Template is missing a secondary example row"
    assert "cca" in kinds, "Template is missing a cca example row"
    print("Template has Secondary + CCA example rows: OK")

    # -- Demo dataset upload, then check Casey Kim's Portfolio. --
    file_input = page.locator('input[type="file"]').first
    file_input.set_input_files(demo_path)
    page.wait_for_timeout(2500)
    page.screenshot(path=f"{OUT}/1_bulk_setup_preview.png", full_page=True)

    page.get_by_text("Confirm and import", exact=True).click()
    page.wait_for_timeout(3000)

    page.get_by_role("tab", name="Associates").click()
    page.wait_for_timeout(1200)
    page.screenshot(path=f"{OUT}/2_associates_list.png", full_page=True)

    page.get_by_role("button", name="Casey Kim", exact=True).click()
    page.wait_for_timeout(1500)
    page.screenshot(path=f"{OUT}/3_casey_portfolio_profile.png", full_page=True)

    page.get_by_role("tab", name="Rotation timeline").click()
    page.wait_for_timeout(1000)
    page.screenshot(path=f"{OUT}/4_casey_portfolio_timeline.png", full_page=True)

    # Expand the "N responsibilities" detail if it's collapsed.
    expanders = page.get_by_text("responsibilities", exact=False)
    if expanders.count():
        expanders.first.click()
        page.wait_for_timeout(800)
    page.screenshot(path=f"{OUT}/5_casey_portfolio_responsibilities_expanded.png", full_page=True)

    browser.close()

print("Screenshots saved to", OUT)
