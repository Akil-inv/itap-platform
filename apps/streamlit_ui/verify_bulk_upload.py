"""One-off Playwright verification of the Bulk Setup tab end to end:
download the template, build a fresh workbook with new names, upload it
through the real browser, confirm the import, and check the result shows
up elsewhere in the app (Org Structure).
"""
import os

from openpyxl import Workbook
from playwright.sync_api import sync_playwright

URL = "http://localhost:8770"
OUT_DIR = "/tmp/itap_screens3"
CHROMIUM = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
UPLOAD_PATH = "/tmp/itap_bulk_upload_test.xlsx"

os.makedirs(OUT_DIR, exist_ok=True)

# Build a small fresh workbook: a new Associate + Manager + Assignment with
# goals and criteria, none of which exist in the seeded demo data.
wb = Workbook()
wb.remove(wb.active)
ws = wb.create_sheet("Associates")
ws.append(["name", "email"])
ws.append(["Riya", "riya@example.com"])
ws = wb.create_sheet("Managers")
ws.append(["name", "email"])
ws.append(["Sam", ""])
ws = wb.create_sheet("Assignments")
ws.append(["associate_name", "manager_name", "start_date", "end_date", "goals", "criteria"])
ws.append(["Riya", "Sam", "2026-01-01", "", "Ship the reporting dashboard", "Communication; Ownership"])
wb.save(UPLOAD_PATH)

with sync_playwright() as p:
    browser = p.chromium.launch(executable_path=CHROMIUM)
    page = browser.new_page(viewport={"width": 1400, "height": 1300})
    page.goto(URL)
    page.wait_for_timeout(2000)

    seed_button = page.get_by_role("button", name="Seed demo data")
    if seed_button.count():
        seed_button.click()
        page.wait_for_timeout(1500)

    page.get_by_role("button", name="Priya", exact=True).click()
    page.wait_for_timeout(1500)

    page.get_by_role("tab", name="Bulk Setup").click()
    page.wait_for_timeout(1000)

    # Verify the template download actually produces a file
    with page.expect_download() as dl_info:
        page.get_by_role("button", name="Download template (.xlsx)").click()
    download = dl_info.value
    saved_path = f"{OUT_DIR}/downloaded_template.xlsx"
    download.save_as(saved_path)
    print("Downloaded template size:", os.path.getsize(saved_path), "bytes")

    page.screenshot(path=f"{OUT_DIR}/1_bulk_setup_empty.png", full_page=True)

    # Upload the fresh workbook
    page.locator('input[type="file"]').set_input_files(UPLOAD_PATH)
    page.wait_for_timeout(1500)
    page.screenshot(path=f"{OUT_DIR}/2_bulk_setup_preview.png", full_page=True)

    page.get_by_role("button", name="Confirm and import").click()
    page.wait_for_timeout(1500)
    page.screenshot(path=f"{OUT_DIR}/3_bulk_setup_result.png", full_page=True)

    # Confirm Riya/Sam now show up in Org Structure
    page.get_by_role("tab", name="Org Structure").click()
    page.wait_for_timeout(1500)
    page.screenshot(path=f"{OUT_DIR}/4_org_structure_after_import.png", full_page=True)

    browser.close()

print("Screenshots saved to", OUT_DIR)
